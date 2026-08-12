from openpyxl import load_workbook
import csv

wb = load_workbook(filename = "book2.xlsx")
sheet=wb["Sheet1"]
row=sheet.max_row
print(row)
